import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Конвертирует CSV в XLSX. Первая строка — заголовок.
    Лист "Sheet1". Автоширина колонок (мин. 8 символов).
    """
    p_in = Path(csv_path)
    if not p_in.exists():
        raise FileNotFoundError(f"Файл не найден: {csv_path}")

    with p_in.open(encoding="utf-8") as f:
        rows = list(csv.reader(f))

    if not rows:
        raise ValueError("Пустой CSV")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for row in rows:
        ws.append(row)

    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        
        width = max(max_len + 2, 8)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    p_out = Path(xlsx_path)
    p_out.parent.mkdir(parents=True, exist_ok=True)
    
    wb.save(p_out)